from pathlib import Path
from typing import Any

from openpyxl import load_workbook


class ExcelReader:
    def __init__(
        self,
        file_path: str | Path,
    ) -> None:
        self.file_path = Path(file_path)

        if not self.file_path.exists():
            raise FileNotFoundError(
                f"找不到 Excel：{self.file_path}"
            )

        self.workbook = load_workbook(
            self.file_path,
            read_only=True,
            data_only=True,
        )

    def sheet_names(self) -> list[str]:
        return list(self.workbook.sheetnames)

    def read_rows(
        self,
        sheet_name: str,
    ) -> list[tuple[Any, ...]]:
        worksheet = self.workbook[
            sheet_name
        ]

        return list(
            worksheet.iter_rows(
                values_only=True
            )
        )